"""ATS-safe .docx generation.

Single column, no tables, no text boxes, no icon fonts, contact details in the
body rather than the header. Those constraints are not aesthetic: header content
and layout tables are exactly what resume parsers drop, and a dropped phone
number is an application that cannot be answered.
"""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt


def build_letter_docx(
    *, title: str, company: str, body: str, display_name: str, contact_line: str = ""
) -> bytes:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    head = doc.add_paragraph()
    run = head.add_run(display_name)
    run.bold = True
    run.font.size = Pt(14)

    if contact_line:
        doc.add_paragraph(contact_line)

    doc.add_paragraph("")
    subject = doc.add_paragraph()
    subject_run = subject.add_run(f"{title} — {company}" if company else title)
    subject_run.bold = True

    doc.add_paragraph("")
    for para in body.split("\n\n"):
        text = para.strip()
        if text:
            doc.add_paragraph(text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
